"""
Arquivo: importar_planilha.py
Propósito:
    Ler a planilha INVENTÁRIO CELULARES.xlsx e importar seus dados para o
    banco SQLite da aplicação, com tratamento básico de cabeçalhos e nulos.
"""

from pathlib import Path

from openpyxl import load_workbook

from database import execute, execute_many

PROJECT_DIR = Path(__file__).resolve().parent
PLANILHAS_CANDIDATAS = [
    PROJECT_DIR.parent / "INVENTÁRIO CELULARES.xlsx",
    PROJECT_DIR.parent / "INVENTÁRIO_CELULARES.xlsx",
]

SHEET_TABLE_MAP = {
    "Inventário": "aparelhos",
    "Chips": "chips",
    "Estoque Aparelhos": "estoque",
    "Celulares devolvidos": "devolvidos",
    "Descarte Eletrônico": "descarte",
    "VW_COLABORADOR": "colaboradores",
    "Funcionários": "colaboradores",
    "Grupos WhatsApp": "grupos_whatsapp",
    "CentroCusto": "centros_custo",
}

ABAS_SEM_TABELA_DIRETA = {
    "Pulsus",
    "Transf. Titularidade",
    "DADOS",
}

TABLE_COLUMNS = {
    "aparelhos": [
        "linha",
        "estab",
        "mdm_id",
        "tipo_linha",
        "status",
        "uso",
        "usuario",
        "nivel",
        "nr_cartao",
        "responsavel",
        "cargo",
        "situacao",
        "centro_custo",
        "departamento",
        "internet",
        "grupo",
        "marca",
        "modelo_mdm",
        "modelo",
        "imei_mdm",
        "imei",
        "chip",
        "mac_wifi",
        "serial",
        "conta_google",
        "observacoes",
    ],
    "chips": ["linha", "chip", "status", "observacoes"],
    "estoque": ["unidade", "marca", "modelo", "imei", "centro_custo", "observacoes", "status"],
    "devolvidos": ["modelo", "imei", "antigo_usuario", "unidade", "observacao", "status"],
    "descarte": ["modelo", "imei", "antigo_usuario", "unidade", "observacao", "status", "data_descarte"],
    "colaboradores": [
        "matricula",
        "nome",
        "estabelecimento",
        "cargo",
        "centro_custo",
        "departamento",
        "data_admissao",
        "data_desligamento",
    ],
    "grupos_whatsapp": [
        "grupo",
        "linha",
        "nome_integrante",
        "administrador",
        "unidade",
        "linha_corporativa",
        "responsavel_informar",
    ],
    "centros_custo": ["codigo", "titulo"],
}


def sanitize_value(value):
    """
    Normaliza valores vindos da planilha.

    Parâmetros:
        value (Any): valor lido da célula Excel.

    Retorno:
        Any: valor tratado, convertendo vazios em None e textos sem espaços.
    """
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def is_header_like(row_values, debug=False, row_num=None):
    """
    Detecta linhas de título/cabeçalho que devem ser ignoradas.

    Parâmetros:
        row_values (list): valores crus da linha.
        debug (bool): se True, imprime logs detalhados.
        row_num (int): numero da linha (para logs).

    Retorno:
        bool: True quando a linha parece ser cabeçalho/título.
    """
    import sys
    prefix = f"[IS_HEADER]"
    if row_num:
        prefix = f"[IS_HEADER][linha {row_num}]"

    non_empty = [str(value).strip().lower() for value in row_values if sanitize_value(value) is not None]

    if debug:
        print(f"{prefix} Valores nao vazios: {non_empty[:10]}... (total: {len(non_empty)})", file=sys.stderr)
        print(f"{prefix} Total de campos na linha: {len(row_values)}", file=sys.stderr)

    if not non_empty:
        if debug:
            print(f"{prefix} IGNORADA: todos valores vazios", file=sys.stderr)
        return True

    # PRIMEIRO: verifica se tem numeros (IMEI, MDM ID, etc.)
    # Se tem numeros, NAO eh cabecalho (eh linha de dados)
    has_number = False
    for value in non_empty:
        if any(c.isdigit() for c in value):
            has_number = True
            break

    if has_number:
        if debug:
            print(f"{prefix} NAO eh cabecalho: tem numeros", file=sys.stderr)
        return False

    # Se NAO tem numeros e a linha tem muitos campos vazios, provavelmente eh cabecalho
    total_fields = len(row_values)
    non_empty_count = len(non_empty)
    ratio = non_empty_count / total_fields if total_fields > 0 else 0
    if debug:
        print(f"{prefix} Ratio campos nao vazios: {ratio:.2f} ({non_empty_count}/{total_fields})", file=sys.stderr)

    if ratio < 0.5:
        if debug:
            print(f"{prefix} IGNORADA: ratio < 0.5 ({ratio:.2f})", file=sys.stderr)
        return True

    # Verifica se eh cabecalho conhecido (nomes exatos de abas)
    joined = " ".join(non_empty)
    exact_headers = [
        "inventário principal",
        "inventário",
        "chips",
        "estoque aparelhos",
        "celulares devolvidos",
        "descarte eletrônico",
        "colaboradores",
        "vw_colaborador",
        "funcionários",
        "grupos whatsapp",
        "centros de custo",
        "centrocusto",
        "pulsus",
        "transf. titularidade",
        "dados",
    ]
    if joined in exact_headers:
        if debug:
            print(f"{prefix} IGNORADA: cabecalho conhecido '{joined}'", file=sys.stderr)
        return True

    # Se chegou aqui, eh provavelmente cabecalho (so texto, sem numeros)
    if debug:
        print(f"{prefix} IGNORADA: so texto, sem numeros", file=sys.stderr)
    return True


def import_sheet(sheet, table_name, debug=False):
    """
    Importa uma aba da planilha para a tabela correspondente.

    Parâmetros:
        sheet (Worksheet): aba carregada pelo openpyxl.
        table_name (str): nome da tabela destino no SQLite.
        debug (bool): se True, imprime logs detalhados.

    Retorno:
        int: quantidade de linhas importadas.
    """
    import sys
    columns = TABLE_COLUMNS[table_name]
    rows_to_insert = []
    row_count = 0
    ignored_empty = 0
    ignored_header = 0

    prefix = f"[IMPORT_SHEET][{sheet.title}->{table_name}]"

    for raw_row in sheet.iter_rows(values_only=True):
        row_count += 1
        row_values = list(raw_row[: len(columns)])
        sanitized = [sanitize_value(value) for value in row_values]

        # Ignora linhas totalmente vazias e alguns títulos/cabeçalhos comuns.
        if all(value is None for value in sanitized):
            ignored_empty += 1
            if debug:
                print(f"{prefix} Linha {row_count}: IGNORADA (todos vazios)", file=sys.stderr)
            continue
        if is_header_like(sanitized, debug=debug, row_num=row_count):
            ignored_header += 1
            continue

        if table_name == "grupos_whatsapp":
            # Converte textos comuns do Excel para booleanos simples no SQLite.
            for index in [3, 5]:
                value = sanitized[index]
                if isinstance(value, str):
                    sanitized[index] = 1 if value.strip().lower() in {"sim", "true", "1", "x"} else 0

        rows_to_insert.append(tuple(sanitized))

    if debug:
        print(f"{prefix} === RESUMO DA IMPORTACAO ===", file=sys.stderr)
        print(f"{prefix} Total de linhas lidas: {row_count}", file=sys.stderr)
        print(f"{prefix} Linhas ignoradas (vazias): {ignored_empty}", file=sys.stderr)
        print(f"{prefix} Linhas ignoradas (cabecalho): {ignored_header}", file=sys.stderr)
        print(f"{prefix} Linhas para inserir: {len(rows_to_insert)}", file=sys.stderr)

    if not rows_to_insert:
        return 0

    placeholders = ", ".join(["?"] * len(columns))
    column_sql = ", ".join(columns)
    execute_many(
        f"INSERT INTO {table_name} ({column_sql}) VALUES ({placeholders})",
        rows_to_insert,
    )
    if debug:
        print(f"{prefix} Registros inseridos: {len(rows_to_insert)}", file=sys.stderr)
    return len(rows_to_insert)


def main(planilha_path=None):
    """
    Executa a importação completa da planilha para o banco.

    Parâmetros:
        planilha_path (str | None): caminho opcional da planilha.

    Retorno:
        dict: relatório com quantidades importadas por aba.
    """
    if planilha_path:
        path = Path(planilha_path)
    else:
        path = next((candidate for candidate in PLANILHAS_CANDIDATAS if candidate.exists()), PLANILHAS_CANDIDATAS[0])

    print(f"Carregando planilha: {path}")
    if not path.exists():
        raise FileNotFoundError(f"Planilha nao encontrada: {path}")

    workbook = load_workbook(path, data_only=True)
    print(f"Abas encontradas: {workbook.sheetnames}")

    importacoes = {}
    for sheet_name, table_name in SHEET_TABLE_MAP.items():
        if sheet_name not in workbook.sheetnames:
            print(f"Aba '{sheet_name}' nao encontrada, pulando...")
            continue
        print(f"Importando aba '{sheet_name}' -> tabela '{table_name}'...")
        sheet = workbook[sheet_name]
        count = import_sheet(sheet, table_name, debug=True)
        importacoes[sheet_name] = count
        print(f"  -> {count} registros importados.")

    if not importacoes:
        print("Nenhuma importacao realizada. Verifique os nomes das abas.")
    else:
        print("\n=== RESUMO ===")
        for sheet_name, count in importacoes.items():
            print(f"  {sheet_name}: {count} registros")

    return importacoes


if __name__ == "__main__":
    main()
